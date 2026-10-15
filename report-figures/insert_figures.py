import struct
from pathlib import Path

import docx
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu, Inches

FIGURES = Path(__file__).resolve().parent
DOC = FIGURES.parent / "Grove-Project-Report.docx"

ORDER = [
    "fig-5-1-wire-catalog.png",
    "fig-5-2-wire-reserve.png",
    "fig-5-3-wire-bookings.png",
    "fig-5-4-wire-admin.png",
    "fig-6-1-home.png",
    "fig-6-2-reserve.png",
    "fig-11-1-payment.png",
    "fig-11-2-receipt.png",
    "fig-14-1-admin.png",
    "fig-14-2-services.png",
    "fig-6-1-home.png",
    "fig-17-2-catalog.png",
    "fig-6-2-reserve.png",
    "fig-11-2-receipt.png",
    "fig-17-5-status.png",
    "fig-14-1-admin.png",
    "fig-14-2-services.png",
    "fig-17-8-users.png",
    "fig-17-9-mobile.png",
]

MAX_WIDTH = 6.3
MAX_HEIGHT = 8.4


def png_size(path):
    data = Path(path).read_bytes()
    return struct.unpack(">II", data[16:24])


def fit(path):
    width, height = png_size(path)
    display_width = MAX_WIDTH
    display_height = display_width * height / width
    if display_height > MAX_HEIGHT:
        display_height = MAX_HEIGHT
        display_width = display_height * width / height
    return Inches(display_width)


def main():
    document = docx.Document(DOC)
    placeholders = [table for table in document.tables if "Paste printout here" in table.cell(0, 0).text]
    if len(placeholders) != len(ORDER):
        raise SystemExit(f"Expected {len(ORDER)} placeholders, found {len(placeholders)}")

    for table, name in zip(placeholders, ORDER):
        path = FIGURES / name
        row = table.rows[0]
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Emu(200000)
        cell = table.cell(0, 0)
        for paragraph in cell.paragraphs:
            paragraph.clear()
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=fit(path))

    for paragraph in document.paragraphs:
        text = paragraph.text
        if text.startswith("This page is left blank"):
            paragraph.text = "Captured from the running Grove application."
        elif "are blank so that printed UI images can be pasted" in text:
            paragraph.text = "The figures in this report were captured from the running Grove application."
        elif "Each figure page is a bordered blank" in text:
            paragraph.text = "Each figure below is a screenshot from the running application."
        elif "The figure pages stay empty until the printouts are pasted." in text:
            paragraph.text = text.replace(
                "The figure pages stay empty until the printouts are pasted.",
                "The figure pages show the screens captured for this demonstration.",
            )

    document.save(DOC)
    print("updated", DOC)


if __name__ == "__main__":
    main()
